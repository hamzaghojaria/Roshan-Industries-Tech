"""Inspect original scans and workbook; writes contact sheets, PDF text and inspection reports."""

# Run directly from a working Python installation; see README for inputs and write effects.

from pathlib import Path
import json
import pymupdf as fitz
from PIL import Image, ImageDraw
from openpyxl import load_workbook


def main():
    """Run this maintenance task explicitly; importing the module never writes files."""
    root = Path(__file__).resolve().parents[1]
    out = root / "reports/high-resolution-audit"
    pdf = fitz.open(root.parent / "reference-documents/roshan-high-resolution-source.pdf")
    records = []
    for i, page in enumerate(pdf):
        rec = {
            "page": i + 1,
            "size": list(page.rect),
            "text": page.get_text(),
            "blocks": page.get_text("blocks"),
            "images": page.get_image_info(xrefs=True),
        }
        records.append(rec)
        if not (out / f"page-{i+1:02}.png").exists():
            pix = page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3))
            pix.save(out / f"page-{i+1:02}.png")
    (out / "pdf-inventory.json").write_text(
        json.dumps(
            records, indent=2, default=lambda v: v.hex() if isinstance(v, bytes) else str(v)
        ),
        encoding="utf8",
    )
    (out / "pdf-text.txt").write_text(
        "\n\n".join(f'PAGE {r["page"]}\n{r["text"]}' for r in records), encoding="utf8"
    )
    wb = load_workbook(root / "Roshan-Industries-Product-Catalogue.xlsx")
    print("WORKBOOK", [(s.title, s.max_row, s.max_column) for s in wb])
    print("PDF PAGES", len(pdf))
    for r in records:
        print(
            "PAGE",
            r["page"],
            "IMAGES",
            len(r["images"]),
            "TEXT",
            r["text"][:170].replace("\n", " | "),
        )
    for start in range(0, len(pdf), 6):
        sheet = Image.new("RGB", (1200, 1350), "#dddddd")
        draw = ImageDraw.Draw(sheet)
        for n in range(start, min(start + 6, len(pdf))):
            im = Image.open(out / f"page-{n+1:02}.png")
            im.thumbnail((380, 630))
            x = ((n - start) % 3) * 400
            y = ((n - start) // 3) * 675
            sheet.paste(im, (x + (400 - im.width) // 2, y + 25))
            draw.text((x + 12, y + 4), f"PAGE {n+1}", fill="black")
        sheet.save(out / f"contact-{start+1:02}.jpg")


if __name__ == "__main__":
    main()
