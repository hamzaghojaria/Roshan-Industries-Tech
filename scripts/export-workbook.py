"""Audit source-photo coverage and create the reviewable Excel catalogue."""

from pathlib import Path
from io import BytesIO
import json, hashlib, csv
import argparse
import pypdfium2 as pdfium
from PIL import Image, ImageOps, ImageChops, ImageDraw
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "reports" / "catalogue-audit"
OUT.mkdir(parents=True, exist_ok=True)
# Treat the supplied PDF only as data; record its hash without modifying it.
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--pdf", type=Path, default=Path.home() / "Downloads" / "20 page.pdf"
)
SOURCE = parser.parse_args().pdf
# Generated data begins with a comment; isolate the assignment before parsing JSON.
products = json.loads(
    (ROOT / "products.js")
    .read_text(encoding="utf8")
    .split("window.ROSHAN_PRODUCTS =", 1)[1]
    .strip()
    .rstrip(";")
)
by_id = {p["id"]: p for p in products}
premium_report = ROOT / "reports" / "premium-catalogue-audit.json"
premium_pages = json.loads(premium_report.read_text(encoding="utf8"))["sku_pages"] if premium_report.exists() else {}
pdf = pdfium.PdfDocument(SOURCE)
assert len(pdf) == 20
expected = [
    f"p{page:02}-{slot:02}"
    for page in range(2, 20)
    for slot in range(1, 13 if page < 18 else 5)
]
assert len(expected) == 200 and set(expected) == set(by_id)
assert len(products) == len({p["sku"] for p in products}) == 200
assert len({p["categoryId"] for p in products}) == 15
columns = [(33, 301), (313, 581), (593, 861)]
rows = [(23, 271), (340, 563), (632, 855), (924, 1147)]
special = {
    18: [
        (38, 28, 430, 419),
        (446, 28, 837, 419),
        (38, 535, 430, 1029),
        (463, 535, 855, 1029),
    ],
    19: [
        (30, 80, 446, 576),
        (462, 80, 869, 576),
        (30, 578, 446, 1190),
        (462, 578, 869, 1190),
    ],
}
# Re-render each source crop and compare it with the corresponding website image.
audit = []
for page in range(2, 20):
    rendered = pdf[page - 1].render(scale=3).to_pil().convert("RGB")
    sheet = Image.new("RGB", (1200, 1050), "#eef2f6")
    draw = ImageDraw.Draw(sheet)
    draw.text(
        (18, 10),
        f"PDF page {page} / Source crop (left) and website image (right)",
        fill="black",
    )
    boxes = special.get(page, [(l, t, r, b) for t, b in rows for l, r in columns])
    for slot, box in enumerate(boxes, 1):
        ident = f"p{page:02}-{slot:02}"
        p = by_id[ident]
        image_path = ROOT / "assets/products" / f"{ident}.webp"
        with Image.open(image_path) as saved:
            assert saved.size == (480, 480)
            actual = saved.convert("RGB")
        crop = rendered.crop(tuple(v * 2 for v in box))
        difference = ImageChops.difference(
            crop, Image.new("RGB", crop.size, "white")
        ).convert("L")
        bounds = difference.point(lambda v: 255 if v > 40 else 0).getbbox()
        assert bounds, ident
        trimmed = crop.crop(bounds)
        ref = Image.new("RGB", (480, 480), "white")
        fit = ImageOps.contain(trimmed, (420, 420), Image.Resampling.LANCZOS)
        ref.paste(fit, ((480 - fit.width) // 2, (480 - fit.height) // 2))
        error = (
            sum(ImageChops.difference(ref, actual).resize((1, 1)).getpixel((0, 0))) / 3
        )
        assert error < 5, (ident, error)
        assert (ROOT / p["url"]).exists() and (ROOT / p["image"]).exists()
        x = ((slot - 1) % 3) * 400
        y = 40 + ((slot - 1) // 3) * 250
        for offset, img in [(0, trimmed), (195, actual)]:
            thumb = ImageOps.contain(img, (190, 210))
            sheet.paste(
                thumb,
                (x + offset + (190 - thumb.width) // 2, y + (210 - thumb.height) // 2),
            )
        draw.text(
            (x + 8, y + 214),
            f'{ident} -> {p["sku"]}',
            fill="black",
        )
        audit.append(
            {
                "source_image": ident,
                "page": page,
                "slot": slot,
                "sku": p["sku"],
                "coverage": "Included",
                "image_comparison_error": round(error, 3),
                "image_sha256": hashlib.sha256(image_path.read_bytes()).hexdigest(),
            }
        )
    sheet.save(OUT / f"page-{page:02}-comparison.jpg", quality=90)

# Shared workbook palette, overview and per-category filtered tables.
navy = "14263D"
blue = "245BDB"
muted = "5B6B7D"
white = "FFFFFF"
wb = Workbook()
summary = wb.active
summary.title = "Overview"
summary.sheet_view.showGridLines = False
summary.merge_cells("A1:F2")
summary["A1"] = "ROSHAN INDUSTRIES | PRODUCT CATALOGUE"
summary["A1"].font = Font(size=22, bold=True, color=white)
summary["A1"].fill = PatternFill("solid", fgColor=navy)
summary["A4"] = "200 products • 15 categories • 200 PDF photo panels"
summary["A4"].font = Font(size=16, bold=True, color=blue)
notes = [
    "Source: 20 page.pdf | Audit: 05 October 2026 | Mumbai, India",
    "All 200 product-photo panels have separate listings. Page 5 selectors have distinct centres.",
    "Pages 1 and 20 are covers, not additional product entries.",
    "Six photos have no printed caption; their website names were confirmed on 05 October 2026.",
    "Other names are transcribed/normalised from the scan, not independently certified specifications.",
    "Page 18 tray and Hands names confirmed; materials and specifications still require confirmation.",
    "Each category sheet includes embedded photos, SKUs, source references and product-page links.",
    "Website links work when this workbook stays beside the website files. Photos remain embedded.",
    "Contact: roshanindustriestech@gmail.com | +91 98212 16170 | WhatsApp: https://wa.me/919821216170",
]
for row, text in enumerate(notes, 6):
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    summary.cell(row, 1, text).alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[row].height = 30
summary["A16"] = "Category"
summary["B16"] = "Products"
summary["C16"] = "Open sheet"
names = {
    "Eye Loupes & Magnifiers": "Eye Loupes & Magnifiers",
    "Precision Screwdrivers": "Precision Screwdrivers",
    "Case Opening Tools": "Case Opening Tools",
    "Bracelet & Strap Tools": "Bracelet & Strap Tools",
    "Glass & Hand Fitting": "Glass & Hand Fitting",
    "Clock Winding Keys": "Clock Winding Keys",
    "Clock Parts & Accessories": "Clock Parts & Accessories",
    "Jewellery Bench Tools": "Jewellery Bench Tools",
    "Soldering & Helping Hands": "Soldering & Helping Hands",
    "Tweezers": "Tweezers",
    "Oil Cups & Oil Pins": "Oil Cups & Oil Pins",
    "Holders & Stands": "Holders & Stands",
    "Gauges & Selectors": "Gauges & Selectors",
    "Trays, Covers & Storage": "Trays Covers & Storage",
    "Compasses": "Compasses",
}
uncertain = {"p17-01", "p17-02", "p19-01", "p19-02", "p19-03", "p19-04"}
ambiguous = {"p18-01", "p18-02", "p18-04", "p03-10", "p03-11", "p17-03", "p16-11"}


def setup(ws, title, headers, widths):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(size=19, bold=True, color=white)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
    ws.cell(3, 1, "Roshan Industries • PDF-based catalogue • 05 October 2026").font = (
        Font(color=muted, size=11)
    )
    ws.cell(4, 1, "← Overview").hyperlink = "#'Overview'!A1"
    ws.cell(4, 1).style = "Hyperlink"
    ws.append(headers)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[5].height = 28
    ws.freeze_panes = "C6"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:5"


def table(ws, number):
    tab = Table(
        displayName=f"Catalogue{number}",
        ref=f"A5:{get_column_letter(ws.max_column)}{ws.max_row}",
    )
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    for cell in ws[5]:
        cell.font = Font(color=white, bold=True)
        cell.fill = PatternFill("solid", fgColor=blue)
    ws.print_options.horizontalCentered = True


# Every category keeps embedded photos, stable SKUs and traceable source references.
def product_sheet(ws, items, photos):
    setup(
        ws,
        ws.title,
        [
            "Photo",
            "SKU",
            "Product name",
            "Category",
            "Family",
            "PDF page",
            "Panel",
            "Image reference",
            "Review / enquiry notes",
            "Website product page",
            "Product description",
            "Custom manufacturing",
            "Premium PDF page",
            "Premium PDF link",
        ],
        [17, 16, 48, 31, 25, 12, 10, 18, 58, 28, 76, 42, 18, 26],
    )
    for p in sorted(items, key=lambda p: p["sku"]):
        note = p["note"]
        if p["id"] in ambiguous and not p.get("nameConfirmed"):
            note = "Original caption is ambiguous; confirm terminology, material and exact specification before ordering."
        ws.append(
            [
                "",
                p["sku"],
                p["name"],
                p["category"],
                p["family"],
                p["page"],
                p["slot"],
                p["id"],
                note,
                "Open product page",
                p["description"],
                "Discuss custom requirements with our team; feasibility and specifications are confirmed on enquiry.",
                premium_pages.get(p["sku"], ""),
                "Open premium catalogue" if p["sku"] in premium_pages else "",
            ]
        )
        r = ws.max_row
        ws.row_dimensions[r].height = 110
        for cell in ws[r]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(size=11, color=navy)
        ws.cell(r, 2).font = Font(size=12, bold=True, color=blue)
        ws.cell(r, 10).hyperlink = p["url"]
        ws.cell(r, 10).style = "Hyperlink"
        if p["sku"] in premium_pages:
            ws.cell(r, 14).hyperlink = f'assets/roshan-updated-product-catalogue.pdf#page={premium_pages[p["sku"]]}'
            ws.cell(r, 14).style = "Hyperlink"
        if photos:
            photo = Image.open(ROOT / p["image"]).convert("RGB")
            photo.thumbnail((105, 105))
            bio = BytesIO()
            photo.save(bio, format="PNG")
            bio.seek(0)
            img = XLImage(bio)
            ws.add_image(img, f"A{r}")
    table(ws, len(wb.worksheets))


groups = {}
for p in products:
    groups.setdefault(p["category"], []).append(p)
for i, (category, items) in enumerate(groups.items(), 17):
    title = names[category]
    ws = wb.create_sheet(title)
    product_sheet(ws, items, True)
    summary.cell(i, 1, category)
    summary.cell(i, 2, len(items))
    summary.cell(i, 3, "View products").hyperlink = f"#'{title}'!A1"
    summary.cell(i, 3).style = "Hyperlink"
master = wb.create_sheet("All Products")
product_sheet(master, products, False)
audit_ws = wb.create_sheet("PDF Audit")
setup(
    audit_ws,
    "PDF PHOTO COVERAGE",
    [
        "PDF page",
        "Panel",
        "Source image",
        "Represented SKU",
        "Coverage",
        "Image comparison error",
    ],
    [14, 12, 22, 22, 58, 27],
)
for a in audit:
    audit_ws.append(
        [
            a["page"],
            a["slot"],
            a["source_image"],
            a["sku"],
            a["coverage"],
            a["image_comparison_error"],
        ]
    )
table(audit_ws, len(wb.worksheets))
audit_ws.freeze_panes = "A6"
for col in "ABCDEF":
    summary.column_dimensions[col].width = {"A": 44, "B": 15, "C": 24}.get(col, 20)
summary.freeze_panes = "A17"
for cell in summary[16]:
    cell.font = Font(bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=blue)
summary["A34"] = "Photo coverage: 200 / 200"
summary["A35"] = "Unique products: 200 / 200"
summary["A36"] = "Uncaptioned photos: 6"
summary["A37"] = "Product names and source references are retained for catalogue traceability."
summary["A39"] = "All Products"
summary["A39"].hyperlink = "#'All Products'!A1"
summary["A39"].style = "Hyperlink"
summary["C39"] = "PDF Audit"
summary["C39"].hyperlink = "#'PDF Audit'!A1"
summary["C39"].style = "Hyperlink"
wb.properties.title = "Roshan Industries Product Catalogue"
wb.properties.creator = "Roshan Industries"
destination = ROOT / "Roshan-Industries-Product-Catalogue.xlsx"
try:
    wb.save(destination)
except PermissionError:
    # Preserve an open workbook; write the requested revision to a new file.
    destination = ROOT / 'Roshan-Industries-Product-Catalogue-Updated.xlsx'
    wb.save(destination)
# Reopen the actual saved file to verify embedded photos, data and local links.
verified = load_workbook(destination)
category_sheets = [verified[names[c]] for c in groups]
for ws in category_sheets + [verified['All Products']]:
    assert 'Name verification' not in [cell.value for cell in ws[5]]
    assert ws.max_column == 14
skus = [ws.cell(r, 2).value for ws in category_sheets for r in range(6, ws.max_row + 1)]
assert len(skus) == len(set(skus)) == 200 and set(skus) == {p["sku"] for p in products}
assert sum(len(ws._images) for ws in category_sheets) == 200
assert verified["PDF Audit"].max_row - 5 == 200
for ws in category_sheets:
    for r in range(6, ws.max_row + 1):
        assert (ROOT / ws.cell(r, 10).hyperlink.target).is_file()
        product = next(p for p in products if p["sku"] == ws.cell(r, 2).value)
        assert ws.cell(r, 3).value == product["name"]
        assert ws.cell(r, 11).value == product["description"]
        if product["sku"] in premium_pages:
            assert ws.cell(r, 13).value == premium_pages[product["sku"]]
            assert ws.cell(r, 14).hyperlink.target.endswith(f'#page={premium_pages[product["sku"]]}')
manifest = {
    "source_pdf_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "source_pages": 20,
    "photo_panels": 200,
    "unique_products": 200,
    "categories": 15,
    "uncaptioned_photos": sorted(uncertain),
    "ambiguous_names": sorted(ambiguous),
    "confirmed_names": [p["id"] for p in products if p.get("nameConfirmed")],
    "name_confirmation_date": "2026-10-05",
    "covers": [1, 20],
    "photos": audit,
}
(OUT / "coverage.json").write_text(json.dumps(manifest, indent=2), encoding="utf8")
with (OUT / "photo-coverage.csv").open("w", newline="", encoding="utf-8-sig") as file:
    writer = csv.DictWriter(file, fieldnames=list(audit[0]))
    writer.writeheader()
    writer.writerows(audit)
(ROOT / "reports" / "Catalogue-Audit.md").write_text(
    """# Roshan Industries catalogue audit\n\nAudited 05 October 2026 against the original `20 page.pdf`.\n\n- 20 source pages: front and back covers plus 18 product pages.\n- 200 photo panels: all have a website image and an assigned listing.\n- 200 unique listings and stable SKUs in 15 categories.\n- Page 5 panels 3 and 4 are distinct selectors: slotted centre RIT-0039 and solid centre RIT-0200.\n- Six uncaptioned images: page 17 panels 1–2 and all four page 19 compasses. Names confirmed by the company on 05 October 2026.\n- Ambiguous captions are flagged in Excel, including page 18 trays and “Hands”; tray material is not asserted.\n- Multi-part sets count as one source photo panel, not one entry per piece.\n\nEach original PDF crop was independently rendered and compared numerically with its website image (mean RGB difference below 5/255). Page-by-page visual comparisons are in `catalogue-audit`. This verifies extraction fidelity; source scan resolution limits available detail. Product terminology, specifications, availability and prices still require company confirmation.\n\nThe workbook was reopened and checked for 200 unique SKUs, all 15 category sheets, 200 embedded category images, 200 audit rows and working local product-page targets. Original PDF SHA-256 and per-image hashes are recorded in `catalogue-audit/coverage.json`.\n""",
    encoding="utf8",
)
print(
    json.dumps(
        {
            "workbook": str(destination),
            "products": 200,
            "category_sheets": 15,
            "embedded_photos": 200,
            "source_panels_verified": 200,
            "largest_image_error": max(a["image_comparison_error"] for a in audit),
            "bytes": destination.stat().st_size,
        }
    )
)
