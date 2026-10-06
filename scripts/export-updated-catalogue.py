"""Rebuild verified exports, then add the photographed pump enquiry range."""
from pathlib import Path
from io import BytesIO
from copy import copy
import json, runpy, hashlib
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Link
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import white, HexColor
from reportlab.platypus import Paragraph
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
# Preserve the existing source-photo audit before extending the customer exports.
base = runpy.run_path(str(ROOT / 'scripts/prepare-premium-pdf.py'))
products = json.loads((ROOT / 'products.js').read_text(encoding='utf8').split('window.ROSHAN_PRODUCTS =')[1].strip().rstrip(';'))
pumps = [p for p in products if p.get('onlineRange')]
groups = {}
for p in pumps:
    groups.setdefault(p['category'], []).append(p)
# Keep each pump category together in the PDF without changing permanent SKUs.
pumps = [p for items in groups.values() for p in items]
pdf_path = base['OUTPUT']
reader = PdfReader(pdf_path)
W, H = base['W'], base['H']
first = len(reader.pages)
writer = PdfWriter()
writer.append(reader)
style = ParagraphStyle('PumpBody', fontName='PremiumBody', fontSize=10, leading=15, textColor=base['navy'])

def paragraph(c, text, x, y, width, size=10):
    para = Paragraph(text.replace('&', '&amp;'), ParagraphStyle('Body', parent=style, fontSize=size, leading=size*1.5))
    _, height = para.wrap(width, H)
    para.drawOn(c, x, y-height)
    return y-height

def heading(c, title):
    # Match the original premium catalogue's logo, rules and category heading.
    c.setFillColor(base['blue'])
    c.rect(0, H-58, 6, 58, fill=1, stroke=0)
    c.drawImage(ImageReader(str(ROOT/'assets/roshan-logo.png')), 38, H-42, 48, 31, preserveAspectRatio=True, mask='auto')
    base['label'](c, 'ROSHAN INDUSTRIES', 96, H-24, 8, bold=True)
    base['label'](c, 'WATCH PARTS & CUSTOM MANUFACTURING', 96, H-38, 5.5, base['muted'])
    c.setStrokeColor(base['line'])
    c.line(38, H-53, W-38, H-53)
    c.line(38, H-70, W-38, H-70)
    base['label'](c, title, 38, H-103, 23, bold=True)

buf = BytesIO()
c = canvas.Canvas(buf, pagesize=(W, H))
pump_annotations = []
heading(c, 'Explore pumps.')
paragraph(c, 'Pumps / 6 categories / 12 products', 38, H-119, W-76, 10)
y = H-170
category_targets = {}
for category, items in groups.items():
    target = first + 1 + len(category_targets)
    category_targets[category] = target
    base['label'](c, category, 51, y, 9, bold=True)
    base['label'](c, f'{len(items)} products', W-147, y, 7, base['muted'])
    base['label'](c, 'VIEW >', W-83, y, 7, base['blue'], True)
    c.setStrokeColor(base['line'])
    c.line(51, y-9, W-51, y-9)
    y -= 26
base['footer'](c, first+1, index=True)
c.showPage()
new_sku_pages = {}
for category, items in groups.items():
    page_index = category_targets[category]
    heading(c, category)
    base['label'](c, f'Pumps / {len(items)} products / Category page 1 of 1', 38, H-124, 10, base['muted'])
    # Reuse the original two-column, three-row category card geometry.
    card_width, card_height = (W-94)/2, 184
    for j, p in enumerate(items):
        x = 38 + (j % 2)*(card_width+18)
        top = H-158-(j//2)*(card_height+19)
        bottom = top-card_height
        c.setFillColor(white)
        c.setStrokeColor(HexColor('#d6e1f6'))
        c.roundRect(x, bottom, card_width, card_height, 8, fill=1, stroke=1)
        c.drawImage(ImageReader(str(ROOT/p['image'])), x+12, top-107, card_width-24, 95, preserveAspectRatio=True, anchor='c', mask='auto')
        base['label'](c, p['sku'], x+12, top-121, 9, base['blue'], True)
        name = Paragraph(p['name'].replace('&','&amp;'), ParagraphStyle('CardName', parent=style, fontName='PremiumBold', fontSize=10, leading=14))
        _, height = name.wrap(card_width-24, 48)
        name.drawOn(c, x+12, top-132-height)
        new_sku_pages[p['sku']] = page_index+1
        pump_annotations.append((page_index,(x,bottom,x+card_width,top),base['website']+'/'+p['url']))
    base['label'](c, 'Enquiry range: specifications and availability confirmed on enquiry.', 38, 75, 7, base['muted'])
    base['label'](c, 'Reference photographs / Photo credits', 38, 62, 6, base['blue'])
    pump_annotations.append((page_index,(38,59,260,71),base['website']+'/photo-credits.html'))
    base['footer'](c, page_index+1, index=True)
    c.showPage()
c.save()
extra = PdfReader(buf)
for page in extra.pages:
    writer.add_page(page)
# Connect the original category index to the pump section without disturbing its links.
overlay_buf = BytesIO()
overlay = canvas.Canvas(overlay_buf, pagesize=(W,H))
base['label'](overlay, 'PUMPS / 6 CATEGORIES / 12 ENQUIRY ENTRIES / VIEW >', 40, 66, 8, base['blue'], True)
overlay.save()
writer.pages[2].merge_page(PdfReader(overlay_buf).pages[0])
writer.add_annotation(2, Link(rect=(40,54,W-40,82), target_page_index=first))
parent = writer.add_outline_item('Pumps enquiry range', first)
for i, (category, target) in enumerate(category_targets.items()):
    y = H-170-26*i
    writer.add_annotation(first, Link(rect=(40,y-12,W-40,y+18), target_page_index=target))
    writer.add_outline_item(category, target, parent=parent)
report = json.loads((ROOT/'reports/premium-catalogue-audit.json').read_text())
report['sku_pages'].update(new_sku_pages)
for number, rect, url in pump_annotations:
    writer.add_annotation(number, Link(rect=rect, url=url))
for number in range(first, len(writer.pages)):
    writer.add_annotation(number, Link(rect=(W-169,15,W-78,37), target_page_index=2))
    writer.add_annotation(number, Link(rect=(225,9,411,23), url='https://wa.me/919821216170'))
writer.add_metadata({'/Title': 'Roshan Industries | 200+ Products', '/Author': 'Roshan Industries', '/Subject': '200 verified catalogue products plus 12 pump enquiry entries'})
with pdf_path.open('wb') as out:
    writer.write(out)
final = PdfReader(pdf_path)
content = '\n'.join(page.extract_text() or '' for page in final.pages)
assert len(final.pages) == first+1+len(groups)
for p in products:
    assert content.count(p['sku']) == 1, p['sku']
internal_links = 0
for page in final.pages:
    for ref in page.get('/Annots', []):
        annotation = ref.get_object()
        dest = annotation.get('/Dest') or annotation.get('/A', {}).get('/D')
        if isinstance(dest, list):
            assert 0 <= final.get_page_number(dest[0].get_object()) < len(final.pages)
            internal_links += 1
(ROOT/'assets/roshan-updated-product-catalogue.pdf').write_bytes(pdf_path.read_bytes())
report.update(pages=len(final.pages), products=len(products), verified_source_products=200, pump_enquiry_entries=len(pumps), categories=21, family_links=5, category_links=21, verified_internal_links=internal_links, verified_pdf_sha256=hashlib.sha256(pdf_path.read_bytes()).hexdigest())
report['category_pages'].update({k:v+1 for k,v in category_targets.items()})
(ROOT/'reports/premium-catalogue-audit.json').write_text(json.dumps(report, indent=2), encoding='utf8')

# Base workbook still independently audits the 200 original source panels.
workbook = runpy.run_path(str(ROOT/'scripts/export-workbook.py'))
destination = workbook['destination']
wb = load_workbook(destination)
master = wb['All Products']
template = wb[workbook['names'][next(iter(workbook['groups']))]]

def add_entry(ws, p, photo=False):
    ws.append(['', p['sku'], p['name'], p['category'], p['family'], '', '', p['id'], p['note'], 'Open product page', p['description'], 'Discuss requirements and feasibility with our team.', report['sku_pages'][p['sku']], 'Open premium catalogue'])
    row = ws.max_row
    for col in range(1,15):
        ws.cell(row,col)._style = copy(template.cell(6,col)._style)
    ws.row_dimensions[row].height = 110 if photo else 60
    ws.cell(row,10).hyperlink = p['url']
    ws.cell(row,14).hyperlink = 'assets/roshan-updated-product-catalogue.pdf#page='+str(report['sku_pages'][p['sku']])
    if photo:
        im = Image.open(ROOT/p['image']).convert('RGBA')
        bg = Image.new('RGBA', im.size, 'white')
        im = Image.alpha_composite(bg, im).convert('RGB')
        im.thumbnail((105,105))
        stream = BytesIO()
        im.save(stream,format='PNG'); stream.seek(0)
        ws.add_image(ExcelImage(stream),f'A{row}')

for category, items in groups.items():
    ws = wb.create_sheet(category)
    workbook['setup'](ws, category, [cell.value for cell in master[5]], [17,16,48,31,25,12,10,18,58,28,76,42,18,26])
    ws.cell(3,1,'Roshan Industries / Pump enquiry range / 06 October 2026')
    for p in items:
        add_entry(ws,p,True)
        add_entry(master,p)
    workbook['table'](ws,len(wb.worksheets))
for table in master.tables.values():
    table.ref=f'A5:N{master.max_row}'
summary=wb['Overview']
summary['A4']='200+ products / 21 categories / 200 verified source products + 12 pump enquiry entries'
for i, (category, items) in enumerate(groups.items(),43):
    summary.cell(i,1,category); summary.cell(i,2,len(items))
    summary.cell(i,3,'View products').hyperlink=f"#'{category}'!A1"
    summary.cell(i,3).style='Hyperlink'
# Internal audit reports and image attribution remain separate from the customer workbook.
for sheet_name in list(wb.sheetnames):
    if 'audit' in sheet_name.lower() or sheet_name == 'Pump Photo Credits':
        del wb[sheet_name]
summary['A4']='200+ products / 21 categories'
summary['A6']='Roshan Industries / Since 1900 / Mumbai, India'
summary['A7']='Browse 212 products across 21 category sheets.'
summary['A8']='Each product has a permanent SKU for easy enquiries.'
summary['A9']='Category sheets include product photographs and descriptions.'
summary['A10']='Contact our team to confirm specifications and availability.'
summary['A11']='Custom manufacturing: share your drawing, sample or requirements.'
summary['A12']='Use the product and PDF links to explore each catalogue entry.'
for row in [34,35,36,37]:
    summary.cell(row,1).value=None
summary['C39'].value=None
summary['C39'].hyperlink=None
wb.save(destination)
check=load_workbook(destination)
assert not any('audit' in name.lower() or 'source' in name.lower() or 'credits' in name.lower() for name in check.sheetnames)
assert check['All Products'].max_row-5 == len(products)
assert sum(len(check[c]._images) for c in groups) == len(pumps)
assert all(p['sku'] in content for p in pumps)
for category in groups:
    for row in check[category].iter_rows(min_row=6):
        assert row[12].value == report['sku_pages'][row[1].value]
        assert (ROOT/row[9].hyperlink.target).is_file()
print(json.dumps({'pdf_pages':len(final.pages),'entries':len(products),'pump_photos':len(pumps),'xlsx':str(destination)}))
(ROOT/'reports/export-audit.json').write_text(json.dumps({'pdf_pages':len(final.pages),'entries':len(products),'categories':21,'embedded_category_photos':212,'verified_source_panels':200,'pump_reference_photos':12,'xlsx':destination.name},indent=2),encoding='utf8')
