"""Export the reviewed website records to a synchronized, styled XLSX workbook."""

# Run directly from a working Python installation; see README for inputs and write effects.

from pathlib import Path
import json
import hashlib
from io import BytesIO
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from workbook_images import preview_bytes, share_image_resources, WORKBOOK_IMAGE_PROFILE

ROOT = Path(__file__).resolve().parents[1]
thumbnail_cache = {}
products = json.loads(
    (ROOT / "products.js")
    .read_text(encoding="utf8")
    .split("window.ROSHAN_PRODUCTS =", 1)[1]
    .strip()
    .rstrip(";")
)
categories = json.loads((ROOT / "src/data/reviewed-categories.json").read_text(encoding="utf8"))
source = ROOT / "assets/roshan-industries-catalogue.pdf"
source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
wb = Workbook()
overview = wb.active
overview.title = "Overview"
navy = "14263D"
blue = "315BD6"
light = "EDF3FF"
muted = "61717C"
family_colors = {
    "Watchmaking": "315BD6",
    "Clockmaking": "B68A42",
    "Jewellery": "8B5CB2",
    "Workshop Essentials": "2C8790",
    "Precision Machining": "4B7F52",
}
headers = [
    "SKU",
    "Product Name",
    "Description",
    "Category",
    "Family",
    "Catalogue Page",
    "Source PDF Panel",
    "Image Reference",
    "Product Page",
    "New Arrival",
    "Photo",
    "Catalogue",
    "Image Width",
    "Image Height",
    "Catalogue SHA256",
]
fields = [
    "sku",
    "name",
    "description",
    "category",
    "family",
    "cataloguePage",
    "slot",
    "image",
    "url",
]


def make_sheet(title, items, index):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    family = (
        items[0]["family"]
        if len(set(p["family"] for p in items)) == 1
        else "All five product families"
    )
    color = family_colors.get(family, navy)
    ws.sheet_properties.tabColor = color
    ws.merge_cells("A1:O1")
    ws["A1"] = "Roshan Industries Product Catalogue"
    ws["A1"].font = Font(name="Calibri", size=19, bold=True, color=navy)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:O2")
    ws["A2"] = f"Family: {family}   |   {title}   |   {len(items)} products"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color=color)
    ws["A3"] = "Back to Overview"
    ws["A3"].hyperlink = "#'Overview'!A1"
    ws["A3"].font = Font(color=blue, underline="single")
    for col, label in enumerate(headers, 1):
        cell = ws.cell(5, col, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 28
    for row, p in enumerate(items, 6):
        values = [p[k] for k in fields] + [
            "Yes" if p["sku"] in selected else "No",
            "",
            p["sourceCatalogue"],
            p["imageWidth"],
            p["imageHeight"],
            source_hash,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.font = Font(name="Calibri", size=10, color=navy)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7FB")
        ws.cell(row, 8).hyperlink = p["image"]
        ws.cell(row, 9).hyperlink = p["url"]
        ws.cell(row, 12).hyperlink = "assets/" + p["sourceCatalogue"]
        ws.cell(row, 5).font = Font(
            name="Calibri", size=10, bold=True, color=family_colors[p["family"]]
        )
        ws.cell(row, 6).hyperlink = (
            f"assets/roshan-industries-catalogue.pdf#page={p['cataloguePage']}"
        )
        for col in [8, 9, 12]:
            ws.cell(row, col).font = Font(name="Calibri", size=10, color=blue, underline="single")
        # Reuse the same encoded preview across All Products and category sheets.
        if p["image"] not in thumbnail_cache:
            thumbnail_cache[p["image"]] = preview_bytes(ROOT / p["image"])
        buffer = BytesIO(thumbnail_cache[p["image"]])
        image = ExcelImage(buffer)
        scale = min(115 / image.width, 100 / image.height)
        image.width *= scale
        image.height *= scale
        ws.add_image(image, f"K{row}")
        ws.row_dimensions[row].height = 88
    widths = [15, 40, 70, 32, 23, 12, 12, 55, 32, 14, 18, 40, 14, 14, 30]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.column_dimensions["O"].hidden = True
    # The Excel table owns its filter. A second worksheet filter on the same range
    # causes desktop Excel to report corrupt content.
    ws.freeze_panes = "C6"
    table = Table(displayName=f"Products{index}", ref=f"A5:O{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:5"
    return ws


selected = set(json.loads((ROOT / "src/data/new-arrivals.json").read_text(encoding="utf8"))["skus"])
overview.sheet_view.showGridLines = False
overview.merge_cells("A1:D1")
overview["A1"] = "Roshan Industries"
overview["A1"].font = Font(size=25, bold=True, color=navy)
overview.row_dimensions[1].height = 38
overview.merge_cells("A2:D2")
overview["A2"] = "Product Catalogue and Database"
overview["A2"].font = Font(size=16, color=blue)
overview.sheet_properties.tabColor = blue
for row, text in enumerate(
    [
        f"{len(products)} products",
        f"{len(categories)} categories across {len(family_colors)} families",
        f"216 source catalogue products and {sum(bool(p.get('onlineProduct')) for p in products)} new online products",
        "Updated website catalogue with clickable category index",
    ],
    4,
):
    overview.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    overview.cell(row, 1, text).font = Font(size=12, color=navy)
overview["A9"] = "All Products"
overview["A9"].hyperlink = "#'All Products'!A1"
overview["A9"].font = Font(size=13, bold=True, color=blue, underline="single")
for col, label in enumerate(["Category", "Family", "Products", "Open Category"], 1):
    overview.cell(11, col, label).fill = PatternFill("solid", fgColor=navy)
    overview.cell(11, col).font = Font(bold=True, color="FFFFFF")
make_sheet("All Products", products, 0)
row = 11
index = 0
for family, color in family_colors.items():
    row += 1
    overview.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    overview.cell(
        row, 1, f"{family.upper()}   {sum(p['family']==family for p in products)} products"
    ).font = Font(size=12, bold=True, color="FFFFFF")
    for col in range(1, 5):
        overview.cell(row, col).fill = PatternFill("solid", fgColor=color)
    overview.row_dimensions[row].height = 30
    for c in [c for c in categories if c["family"] == family]:
        index += 1
        title = c["name"].replace(" and ", " ")[:31]
        items = [p for p in products if p["categoryId"] == c["id"]]
        make_sheet(title, items, index)
        row += 1
        for col, value in enumerate([c["name"], c["family"], len(items), "View Products"], 1):
            overview.cell(row, col, value).alignment = Alignment(wrap_text=True, vertical="center")
        overview.cell(row, 2).font = Font(bold=True, color=color)
        overview.cell(row, 4).hyperlink = f"#'{title}'!A1"
        overview.cell(row, 4).font = Font(color=blue, underline="single")
        overview.row_dimensions[row].height = 28
for col, width in zip("ABCD", [40, 25, 14, 22]):
    overview.column_dimensions[col].width = width
overview.freeze_panes = "A12"
output = ROOT / "Roshan-Industries-Product-Catalogue.xlsx"
wb.save(output)
image_resources = share_image_resources(output)
check = load_workbook(ROOT / "Roshan-Industries-Product-Catalogue.xlsx")
ws = check["All Products"]
rows = list(ws.iter_rows(min_row=6, values_only=True))
assert len(rows) == len(products)
for values, p in zip(rows, products):
    assert list(values[:9]) == [p[k] for k in fields]
    assert values[9] == ("Yes" if p["sku"] in selected else "No")
assert len(ws._images) == len(products)
assert sum(len(s._images) for s in check) == len(products) * 2
assert all(s.auto_filter.ref is None for s in check)
assert all(len(s.tables) == 1 for s in list(check)[1:])
assert all(s.sheet_properties.tabColor is not None for s in check)
assert all("Family:" in str(s["A2"].value) for s in list(check)[1:])
(ROOT / "reports/high-resolution-audit/workbook-validation.json").write_text(
    json.dumps(
        {
            "products": len(rows),
            "categorySheets": len(categories),
            "embeddedPhotos": len(products) * 2,
            "imageOptimization": WORKBOOK_IMAGE_PROFILE,
            "xlsxBytes": output.stat().st_size,
            **image_resources,
            "sourceSha256": source_hash,
            "familyColors": family_colors,
            "allTabsColored": True,
            "allFieldsMatchWebsite": True,
            "newArrivalProducts": len(selected),
            "xlsxSha256": hashlib.sha256(
                (ROOT / "Roshan-Industries-Product-Catalogue.xlsx").read_bytes()
            ).hexdigest(),
        },
        indent=2,
    ),
    encoding="utf8",
)
print(
    f"PASS XLSX: {len(products)} synchronized records, {len(categories)} category sheets, {len(products)*2} embedded photos, colored family tabs and matching PDF references.",
    flush=True,
)
