"""Create a compact, clickable opening sheet without changing catalogue entries."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def simplify_overview(wb):
    """Keep only branding, totals, category navigation and two contact links."""
    if 'Overview' in wb.sheetnames:
        del wb['Overview']
    ws = wb.create_sheet('Overview', 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    navy, blue, muted = '152A35', '315BD6', '61717C'
    for column, width in [('A', 48), ('B', 12), ('C', 18)]:
        ws.column_dimensions[column].width = width
    ws.merge_cells('A1:C2')
    ws['A1'] = 'ROSHAN INDUSTRIES'
    ws['A1'].font = Font(name='Segoe UI', size=23, bold=True, color=navy)
    ws['A1'].alignment = Alignment(vertical='center')
    ws.merge_cells('A3:C3')
    ws['A3'] = 'Product catalogue | Since 1900'
    ws['A3'].font = Font(name='Segoe UI', size=11, color=muted)
    categories = [sheet for sheet in wb if sheet.title not in ['Overview', 'All Products']]
    total = wb['All Products'].max_row - 5
    ws.merge_cells('A4:C4')
    ws['A4'] = f'200+ products | {len(categories)} categories'
    ws['A4'].font = Font(name='Segoe UI', size=12, bold=True, color=blue)
    ws.append([])
    # Fixed row positions keep the first category visible without scrolling.
    ws.cell(6, 1, 'All products').hyperlink = "#'All Products'!A1"
    ws.cell(6, 2, total)
    ws.cell(6, 3, 'Open catalogue').hyperlink = "#'All Products'!A1"
    for column, label in enumerate(['Category', 'Products', 'Open sheet'], 1):
        cell = ws.cell(7, column, label)
        cell.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=navy)
    for row, sheet in enumerate(categories, 8):
        target = "#'" + sheet.title.replace("'", "''") + "'!A1"
        ws.cell(row, 1, sheet.title).hyperlink = target
        ws.cell(row, 2, sheet.max_row - 5)
        ws.cell(row, 3, 'View products').hyperlink = target
        for cell in ws[row]:
            cell.fill = PatternFill('solid', fgColor='F3F6FA' if row % 2 == 0 else 'FFFFFF')
    for row in range(6, 8 + len(categories)):
        ws.row_dimensions[row].height = 25
        for cell in ws[row]:
            if row != 7:
                cell.font = Font(name='Segoe UI', size=11, color=blue if cell.hyperlink else navy,
                                 underline='single' if cell.hyperlink else None)
            cell.alignment = Alignment(vertical='center', horizontal='center' if cell.column == 2 else 'left')
    contact_row = 9 + len(categories)
    for row, text, target in [
        (contact_row, 'roshanindustriestech@gmail.com', 'mailto:roshanindustriestech@gmail.com'),
        (contact_row + 1, 'Visit website', 'https://roshan-industries-tech.onrender.com/'),
    ]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = ws.cell(row, 1, text)
        cell.hyperlink = target
        cell.font = Font(name='Segoe UI', size=11, color=blue, underline='single')
        ws.row_dimensions[row].height = 24
    ws.freeze_panes = 'A8'
    ws.print_area = f'A1:C{contact_row + 1}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    for sheet in wb:
        sheet.sheet_view.tabSelected = sheet is ws
    wb.active = 0


if __name__ == '__main__':
    path = Path(__file__).resolve().parent.parent / 'Roshan-Industries-Product-Catalogue.xlsx'
    workbook = load_workbook(path)
    simplify_overview(workbook)
    workbook.save(path)
    verified = load_workbook(path)
    assert verified.active.title == 'Overview'
    assert verified['Overview'].max_row == 31
    assert verified['All Products'].max_row - 5 == 212
    assert sum(len(sheet._images) for sheet in verified) == 212
    for row in verified['Overview'].iter_rows(min_row=8, max_row=28):
        assert row[0].value in verified.sheetnames
        assert row[2].hyperlink.target.startswith("#'")
    print('PASS: compact opening sheet, 21 clickable categories, 212 entries/photos preserved')
